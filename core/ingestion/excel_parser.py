"""Deterministic scenario-spreadsheet ingestion (openpyxl, no LLM).

Each data row becomes:
  * one ``scenarios`` row (upserted, keyed by project + scenario number), and
  * one verbatim ``evidence_vault`` row whose ``source_reference`` pins the
    exact sheet/row range, e.g. ``80_Scenarios.xlsx!Sheet1!A14:F14``.

All rows are parsed and validated before anything is written, and the writes
happen in a single transaction, so a bad row never leaves a half-ingested
sheet behind.
"""

from __future__ import annotations

import hashlib
import os
import re
from pathlib import Path
from typing import Any, BinaryIO, Optional, Union

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core.storage.db import ProjectMemoryDB, normalize_status

# normalised header text -> scenario field
HEADER_ALIASES: dict[str, str] = {
    "scenario": "scenario_num",
    "scenarionum": "scenario_num",
    "scenariono": "scenario_num",
    "scenarionumber": "scenario_num",
    "scenarioid": "scenario_num",
    "title": "title",
    "scenariotitle": "title",
    "description": "description",
    "details": "description",
    "status": "status",
    "blockers": "blockers",
    "blocker": "blockers",
    "linkedrequirement": "linked_req_id",
    "linkedreq": "linked_req_id",
    "linkedreqid": "linked_req_id",
    "requirement": "linked_req_id",
    "requirementid": "linked_req_id",
}
REQUIRED_FIELDS = ("scenario_num", "title")
HEADER_SCAN_ROWS = 10


def _norm_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).lower()) if value is not None else ""


def _cell_text(value: Any) -> str:
    """Render a cell value as stable verbatim text (14.0 -> '14')."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_scenario_num(value: Any, where: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{where}: invalid Scenario # {value!r}")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    match = re.fullmatch(r"\s*(?:[A-Za-z]+[-_ ]?)?0*(\d+)\s*", str(value))
    if not match:
        raise ValueError(f"{where}: invalid Scenario # {value!r}")
    return int(match.group(1))


def _find_header(rows: list[tuple]) -> Optional[tuple[int, dict[int, str], dict[int, str]]]:
    """Return (header_row_index0, {col_idx: field}, {col_idx: original header})."""
    for idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        mapping: dict[int, str] = {}
        labels: dict[int, str] = {}
        for col, value in enumerate(row):
            field = HEADER_ALIASES.get(_norm_header(value))
            if field and field not in mapping.values():
                mapping[col] = field
                labels[col] = _cell_text(value)
        if all(f in mapping.values() for f in REQUIRED_FIELDS):
            return idx, mapping, labels
    return None


def ingest_scenarios_excel(
    db: ProjectMemoryDB,
    project_id: str,
    file_path: Union[str, os.PathLike, BinaryIO],
) -> dict:
    """Ingest every sheet that has a recognisable scenario header row.

    Returns ``{"rows_ingested", "scenario_ids", "evidence_ids"}``.
    """
    db.require_project(project_id)
    if isinstance(file_path, (str, os.PathLike)):
        file_label = Path(file_path).name
    else:
        file_label = Path(getattr(file_path, "name", "") or "in-memory.xlsx").name

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheets = [(ws.title, list(ws.iter_rows(values_only=True))) for ws in workbook.worksheets]
    finally:
        workbook.close()

    parsed: list[tuple[dict, dict]] = []  # (scenario_data, evidence kwargs)
    seen_nums: dict[int, str] = {}
    found_header = False
    for sheet_name, rows in sheets:
        header = _find_header(rows)
        if header is None:
            continue
        found_header = True
        header_idx, mapping, labels = header
        last_col = get_column_letter(max(max(mapping) + 1, max((len(r) for r in rows), default=1)))

        for idx0 in range(header_idx + 1, len(rows)):
            row = rows[idx0]
            excel_row = idx0 + 1
            if all(_cell_text(v) == "" for v in row):
                continue
            where = f"{file_label}!{sheet_name}!{excel_row}"

            values = {field: (row[col] if col < len(row) else None) for col, field in mapping.items()}
            if _cell_text(values.get("scenario_num")) == "":
                raise ValueError(f"{where}: missing Scenario #")
            if _cell_text(values.get("title")) == "":
                raise ValueError(f"{where}: missing Title")

            num = _parse_scenario_num(values["scenario_num"], where)
            if num in seen_nums:
                raise ValueError(f"{where}: duplicate Scenario # {num} (first seen at {seen_nums[num]})")
            seen_nums[num] = where

            try:
                status = normalize_status(values.get("status"))
            except ValueError as exc:
                raise ValueError(f"{where}: {exc}") from None
            # Controlled authority: a spreadsheet cell is a claim, not proof.
            # Without linked GIT_DIFF/TEST_LOG evidence it cannot be 'Verified'.
            if status == "Verified":
                status = "Impl_Reported"

            scenario = {
                "scenario_num": num,
                "title": _cell_text(values["title"]),
                "description": _cell_text(values.get("description")) or None,
                "status": status,
                "blockers": _cell_text(values.get("blockers")) or None,
                "linked_req_id": _cell_text(values.get("linked_req_id")) or None,
            }

            verbatim = " | ".join(
                f"{labels[col]}: {_cell_text(row[col] if col < len(row) else None)}"
                for col in sorted(mapping)
            )
            digest = hashlib.sha256(verbatim.encode("utf-8")).hexdigest()[:12]
            evidence = {
                "evidence_id": f"EXCEL:{file_label}:{sheet_name}!{excel_row}:{digest}",
                "source_type": "EXCEL",
                "source_reference": f"{file_label}!{sheet_name}!A{excel_row}:{last_col}{excel_row}",
                "content": verbatim,
                "byte_range": None,
            }
            parsed.append((scenario, evidence))

    if not found_header:
        raise ValueError(f"{file_label}: no sheet with 'Scenario #' and 'Title' headers found")

    scenario_ids: list[str] = []
    evidence_ids: list[str] = []
    with db.transaction():
        for scenario, evidence in parsed:
            current = db.get_scenario(project_id, scenario["scenario_num"])
            if (
                current is not None
                and current["status"] == "Verified"
                and scenario["status"] == "Impl_Reported"
            ):
                # Never let an unproven sheet cell downgrade evidence-backed status.
                scenario["status"] = "Verified"
            scenario_ids.append(db.upsert_scenario(project_id, scenario))
            evidence_ids.append(db.add_evidence(project_id, **evidence))

    return {
        "rows_ingested": len(parsed),
        "scenario_ids": scenario_ids,
        "evidence_ids": evidence_ids,
    }
